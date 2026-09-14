"""
Ground Data Processing Tool - standalone command-line port.

Faithful offline port of the Power Apps / Dataverse React app
(apps/csv-productivity-builder). It accepts a CSV or Excel upload, applies the
same parse -> derive -> build-productivity pipeline, and writes the resulting
productivity table to CSV or Excel.

Reference data that the original app read from Dataverse (Team Composition,
Municipality Code, Mapping Status, Batangas Build) plus existing records is
supplied here as an optional reference workbook (one sheet per table). When a
table is missing the tool falls back to the same defaults the app used
(team -> NO MATCH, build -> BALANCE, etc.).

Usage:
    productivity_tool --input rows.xlsx --reference reference.xlsx --output out.xlsx
    productivity_tool --input rows.csv  --output out.csv --mode negotiation

Run with no arguments for an interactive prompt (double-click friendly).
"""

from __future__ import annotations

import argparse
import collections
import csv
import math
import os
import re
import sys
from datetime import datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP

try:
    from openpyxl import load_workbook, Workbook
except Exception:  # pragma: no cover - only hit when dependency missing
    load_workbook = None
    Workbook = None


# ---------------------------------------------------------------------------
# Column definitions (ported verbatim from src/pages/index.tsx)
# ---------------------------------------------------------------------------

BATANGAS_NEGO_RECORD_COLUMNS = [
    'ID', 'TEAM', 'GROUP', 'NEGOTIATOR NAME', 'AREA-INDEX', 'INDEX NO', 'TYPE OF REPORT', 'PROVINCE', 'MUNICIPALITY',
    'BARANGAY', 'LOT NO', 'SURVEY NO', 'PROPERTY IDENTIFICATION NO (PIN)', 'REGISTERED OWNER', 'AUTHORIZED REPRESENTATIVE',
    'TITLE TYPE', 'TITLE NO', 'LOT AREA (SQM)', 'CLASSIFICATION', 'CLASS', 'NEGO CODE', 'ACTION', 'NEGO DATE',
    'DATE UPLOADED', 'DATE EXTRACTED', 'NEXT CALL DATE', 'ACTION ITEM', 'NEGO INDICATIVE PRICE - ASKING (SALE)',
    'NEGO INDICATIVE PRICE - OFFER (SALE)', 'NET OR GROSS (SALE)', 'NEGO FIRST PAYMENT', 'NEGO PAYMENT TERMS',
    'NEGO INDICATIVE PRICE - ASKING (LEASE)', 'NEGO INDICATIVE PRICE - OFFER (LEASE)', 'NET OR GROSS (LEASE)',
    'ADVANCE PAYMENT (YRS)', 'NEGO CONTRACT TERMS (LEASE)', 'NEGO LEASE ESCALATION RATE', 'PRE-RCC LEVEL',
    'AVAILABILITY/ SCHEDULING WITH THE LO', 'PENDING TITLE - UNDOCUMENTED LAND', 'PENDING TITLE - UNTITLED LAND',
    'PENDING TITLE - ODC', 'WITH ANNOTATION - ACTIVE ADVERSE CLAIM',
    'WITH ANNOTATION - THIRD-PARTY ENCUMBRANCES/MORTGAGE', 'WITH TENANT ISSUE', 'WITH CLAIMANT ISSUE',
    'WITH OVERLAP ISSUE', 'LANDOWNER/AIF VERIFICATION', 'VERIFICATION OF HEIRS/AIF OF THE DECEASED RO/S',
    'HEIRS/AIF BASED ABROAD', 'NO NEXT OF KIN TO THE RO', 'AWAITING DECISION FROM THE LO', 'XCOORD', 'YCOORD',
    'CORRECT GROUP', 'CORRECT TEAM', 'HEAD NEGOTIATOR', 'PRICE VARIANCE (SALE)',
    'MAPPING STATUS', 'MUNICODE', 'LOT AREA (HA)', 'BUILD', 'DATA USABILITY', 'YEAR',
]

LAND_SOURCING_RECORD_COLUMNS = [
    'ID', 'TEAM', 'GROUP', 'LSA NAME', 'AREA-INDEX', 'INDEX NO', 'REPORT DATE', 'DATE UPLOADED', 'DATE EXTRACTED',
    'TYPE OF REPORT', 'PROVINCE', 'MUNICIPALITY', 'BARANGAY', 'LOT NO', 'SURVEY NO',
    'PROPERTY IDENTIFICATION NO (PIN)', 'REGISTERED OWNER', 'AUTHORIZED REPRESENTATIVE', 'TITLE TYPE', 'TITLE NO',
    'LOT AREA (SQM)', 'CLASSIFICATION', 'CLASS', 'INDICATIVE PRICE SALE', 'PAYMENT TERMS', 'NET OR GROSS',
    'INDICATIVE PRICE LEASE', 'CONTRACT TERMS (LEASE)', 'ESCALATION', 'XCOORD', 'YCOORD', 'MAPPING STATUS',
    'MUNICODE', 'LOT AREA (HA)', 'BUILD', 'DATA USABILITY', 'YEAR',
]

DERIVED_TEMPLATE_COLUMNS = {'ID', 'MAPPING STATUS', 'MUNI CODE', 'MUNICODE', 'LOT AREA (HA)', 'BUILD', 'DATA USABILITY', 'YEAR', 'LO'}
OPTIONAL_UPLOAD_COLUMNS = {'CORRECT GROUP', 'CORRECT TEAM', 'HEAD NEGOTIATOR', 'PRICE VARIANCE (SALE)', 'LO'}
# Derived review columns that calculate_review_rows recomputes — a manual edit to
# one of these is remembered in row['_overrides'] so re-calculating keeps it.
DERIVED_REVIEW_COLUMNS = {'AREA-INDEX', 'MUNICODE', 'MAPPING STATUS', 'BUILD', 'DATA USABILITY', 'YEAR', 'LOT AREA (HA)'}

REQUIRED_COLUMNS_BY_MODE = {
    'negotiation': ['ID', 'NEGOTIATOR NAME', 'AREA-INDEX', 'NEGO DATE', 'CLASSIFICATION', 'MUNICIPALITY', 'DATA USABILITY'],
    'land-sourcing': ['ID', 'LSA NAME', 'AREA-INDEX', 'REPORT DATE', 'CLASSIFICATION', 'MUNICIPALITY', 'DATA USABILITY'],
}

EXPLICITLY_NORMALIZED_DATE_COLUMNS = {'REPORT DATE', 'DATE UPLOADED', 'DATE EXTRACTED'}

# Dates recording work that already happened, so a future value is a typo rather
# than a plan. NEXT CALL DATE is deliberately absent: it is the one date column
# meant to be ahead of today.
NO_FUTURE_DATE_COLUMNS = {'NEGO DATE', 'SOURCING DATE', 'REPORT DATE'}

# Mapping-status key -> label maps (from generated/models/mappingstatus-model.ts)
MAPPING_STATUS_DATA_USABILITY_LABEL = {'DataUsability': 'Data Usability', 'VALID': 'VALID', 'INVALID': 'INVALID'}
MAPPING_STATUS_LABEL = {
    'MappingStatus': 'Mapping Status', 'NEWLYMAPPED': 'NEWLY MAPPED', 'INDEXEDPROPERTYUPDATE': 'INDEXED PROPERTY UPDATE',
    'NEWLYSOURCED': 'NEWLY SOURCED', 'DUPLICATE': 'DUPLICATE', 'INCORRECTGEOTAG': 'INCORRECT GEOTAG',
    'INCOMPLETEINFORMATION': 'INCOMPLETE INFORMATION',
}


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def to_fixed(value: float, digits: int = 4) -> float:
    """Mimic JS Number.prototype.toFixed rounding (half away from zero)."""
    if value is None or (isinstance(value, float) and (math.isnan(value) or math.isinf(value))):
        return 0.0
    quant = Decimal(1).scaleb(-digits)
    return float(Decimal(str(value)).quantize(quant, rounding=ROUND_HALF_UP))


def round_whole_number_text(value) -> str:
    """Port of the workspace roundWholeNumberText (JS Math.round -> string)."""
    try:
        n = float(str(value).replace(',', ''))
    except (ValueError, TypeError):
        return '' if value is None else str(value)
    if not math.isfinite(n):
        return str(value)
    return str(math.floor(n + 0.5))  # JS Math.round rounds half toward +inf


def js_number_str(value) -> str:
    """Render a value the way JS String(Number) does: integral floats drop '.0'."""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return '' if value is None else str(value)


def canonical_header(value: str) -> str:
    """Normalize header to alphanumeric-only uppercase for flexible, robust column matching."""
    s = str(value or '').upper()
    return re.sub(r'[^A-Z0-9]', '', s)


def normalize_header(value: str) -> str:
    value = (value or '').lstrip('\ufeff').strip()
    value = re.sub(r'[–—\ufffd]', '-', value)  # en/em dash / replacement char -> hyphen
    value = re.sub(r'\s+', ' ', value)
    return value.upper()


def normalize_team(value: str | None) -> str:
    s = str(value if value is not None else '').strip()
    if not s or s.upper() in ('UNASSIGNED', '0', 'TEAM 0', 'NONE', 'N/A', 'NULL', 'UNMATCHED'):
        return 'TEAM 0'
    return s


def normalize_group(value: str | None) -> str:
    s = str(value if value is not None else '').strip()
    if not s or s.upper() in ('UNASSIGNED', '0', 'GROUP 0', 'NONE', 'N/A', 'NULL', 'UNMATCHED'):
        return 'GROUP 0'
    return s


def is_valid_mm_dd_yyyy(value: str) -> bool:
    m = re.match(r'^(\d{2})/(\d{2})/(\d{4})$', value or '')
    if not m:
        return False
    month, day, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        d = datetime(year, month, day)
    except ValueError:
        return False
    return d.year == year and d.month == month and d.day == day


_EXCEL_EPOCH = datetime(1899, 12, 30, tzinfo=timezone.utc)


def is_excel_date_serial(value) -> bool:
    s = str(value or '').strip()
    if re.match(r'^\d{4,5}(?:\.\d+)?$', s):
        try:
            val = float(s)
            return math.isfinite(val) and 1 <= val <= 60000
        except (ValueError, TypeError):
            return False
    return False


def excel_serial_to_date_str(value) -> str:
    s = str(value or '').strip()
    try:
        serial = float(s)
        if math.isfinite(serial) and 1 <= serial <= 60000:
            d = _EXCEL_EPOCH + timedelta(days=serial)
            return f"{d.month:02d}/{d.day:02d}/{d.year}"
    except Exception:
        pass
    return s


def date_to_excel_serial(value) -> str:
    """Convert any date representation (e.g. '8/24/2026', '2026-08-24') to Excel serial string (e.g. '46258')."""
    if value is None:
        return ''
    if isinstance(value, datetime):
        epoch = datetime(1899, 12, 30, tzinfo=value.tzinfo if value.tzinfo else timezone.utc)
        return str((value - epoch).days)
    s = str(value).strip()
    if not s:
        return ''
    # Already an Excel serial number?
    if is_excel_date_serial(s):
        try:
            return str(int(float(s)))
        except (ValueError, TypeError):
            return s
    # Parse MM/DD/YYYY or M/D/YYYY
    m = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$', s)
    if m:
        month, day, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            d = datetime(year, month, day, tzinfo=timezone.utc)
            return str((d - _EXCEL_EPOCH).days)
        except Exception:
            pass
    # Parse ISO YYYY-MM-DD
    m_iso = re.match(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})', s)
    if m_iso:
        year, month, day = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
        try:
            d = datetime(year, month, day, tzinfo=timezone.utc)
            return str((d - _EXCEL_EPOCH).days)
        except Exception:
            pass
    # Permissive fallback
    for fmt in ('%m/%d/%Y', '%m/%d/%y', '%B %d, %Y', '%b %d, %Y', '%d %B %Y', '%Y-%m-%dT%H:%M:%S'):
        try:
            d = datetime.strptime(s, fmt).replace(tzinfo=timezone.utc)
            return str((d - _EXCEL_EPOCH).days)
        except Exception:
            continue
    return s



def format_date_to_mm_dd_yyyy(value) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return f"{value.month:02d}/{value.day:02d}/{value.year}"
    trimmed = str(value).strip()
    if not trimmed:
        return ''

    # Excel serial number
    if is_excel_date_serial(trimmed):
        return excel_serial_to_date_str(trimmed)


    # ISO-ish yyyy-mm-dd or yyyy/mm/dd
    m = re.match(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})', trimmed)
    if m:
        year, month, day = m.group(1), m.group(2), m.group(3)
        return f"{int(month):02d}/{int(day):02d}/{year}"

    # slash / dash date m/d/yy(yy) (including timestamps)
    m = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})', trimmed)
    if m:
        first, second, raw_year = m.group(1), m.group(2), m.group(3)
        year = f"20{raw_year}" if len(raw_year) == 2 else raw_year
        formatted = f"{int(first):02d}/{int(second):02d}/{year}"
        return formatted if is_valid_mm_dd_yyyy(formatted) else trimmed

    # permissive fallback parse
    for fmt in ('%m/%d/%Y', '%m/%d/%y', '%B %d, %Y', '%b %d, %Y', '%d %B %Y', '%Y-%m-%dT%H:%M:%S'):
        try:
            d = datetime.strptime(trimmed, fmt)
            return f"{d.month:02d}/{d.day:02d}/{d.year}"
        except ValueError:
            continue
    return trimmed


def is_date_column(column: str) -> bool:
    return 'DATE' in column.upper() or column in EXPLICITLY_NORMALIZED_DATE_COLUMNS


def normalize_csv_date_columns(row: dict) -> dict:
    return {k: (format_date_to_mm_dd_yyyy(v) if is_date_column(k) else v) for k, v in row.items()}


def is_future_date(value: str) -> bool:
    """True when a valid MM/DD/YYYY value falls after today (local date)."""
    m = re.match(r'^(\d{2})/(\d{2})/(\d{4})$', value or '')
    if not m:
        return False
    try:
        parsed = datetime(int(m.group(3)), int(m.group(1)), int(m.group(2)))
    except ValueError:
        return False
    return parsed.date() > datetime.now().date()


def date_cell_error(column: str, value) -> str:
    """Why a date cell is unusable, or '' when it is fine.

    Shared by the pre-generate and pre-publish checks so both enforce one rule.
    """
    raw = str(value or '').strip()
    if not raw:
        return ''
    formatted = format_date_to_mm_dd_yyyy(raw)
    if not is_valid_mm_dd_yyyy(formatted):
        return 'needs MM/DD/YYYY (e.g. 08/10/2026)'
    if column.strip().upper() in NO_FUTURE_DATE_COLUMNS and is_future_date(formatted):
        return 'cannot be a future date'
    return ''


def get_invalid_date_cells(rows: list[dict], columns: list[str]) -> list[dict]:
    """Port of getInvalidDateCells: any non-empty date cell that isn't valid
    MM/DD/YYYY, plus nego/sourcing dates that fall after today."""
    invalid = []
    for index, row in enumerate(rows):
        for column in columns:
            if not is_date_column(column):
                continue
            value = str(row.get(column, '') or '').strip()
            reason = date_cell_error(column, value)
            if reason:
                invalid.append({'row': index + 1, 'column': column,
                                'value': value, 'reason': reason})
    return invalid


# ---------------------------------------------------------------------------
# CSV parsing (ported from parseDelimitedText / getCoalescedParsedCsvRows / parseCsv)
# ---------------------------------------------------------------------------

def parse_delimited_text(text: str) -> list[list[str]]:
    cleaned = text.lstrip('﻿')
    first_line = re.split(r'\r?\n', cleaned)[0] if cleaned else ''
    candidates = [',', '\t', ';']
    delimiter = max(candidates, key=lambda c: first_line.count(c) + 1) if first_line else ','
    # match JS tie-break: highest split-count wins, first candidate on ties
    best = None
    best_count = -1
    for c in candidates:
        count = len(first_line.split(c))
        if count > best_count:
            best_count = count
            best = c
    delimiter = best or ','

    rows: list[list[str]] = []
    current_row: list[str] = []
    current_value = ''
    in_quotes = False
    i = 0
    n = len(cleaned)
    while i < n:
        char = cleaned[i]
        next_char = cleaned[i + 1] if i + 1 < n else ''
        if char == '"' and in_quotes and next_char == '"':
            current_value += '"'
            i += 1
        elif char == '"':
            in_quotes = not in_quotes
        elif char == delimiter and not in_quotes:
            current_row.append(current_value.strip())
            current_value = ''
        elif (char == '\n' or char == '\r') and not in_quotes:
            if char == '\r' and next_char == '\n':
                i += 1
            current_row.append(current_value.strip())
            if any(len(v) > 0 for v in current_row):
                rows.append(current_row)
            current_row = []
            current_value = ''
        else:
            current_value += char
        i += 1

    current_row.append(current_value.strip())
    if any(len(v) > 0 for v in current_row):
        rows.append(current_row)
    return rows


def get_expected_csv_upload_columns(mode: str) -> list[str]:
    record_columns = LAND_SOURCING_RECORD_COLUMNS if mode == 'land-sourcing' else BATANGAS_NEGO_RECORD_COLUMNS
    return [c for c in record_columns if c not in DERIVED_TEMPLATE_COLUMNS and c not in OPTIONAL_UPLOAD_COLUMNS]


def coalesce_parsed_rows(parsed_rows: list[list[str]], mode: str) -> list[list[str]]:
    headers = parsed_rows[0] if parsed_rows else []
    # Merge wrapped continuation lines only up to the file's own column count. Using
    # the full expected width here would wrongly collapse every row of a file that is
    # simply missing one or more columns (handled by add_missing_upload_columns) into
    # a single row, since those rows are shorter than the expected width.
    expected = len(headers)
    data_rows: list[list[str]] = []
    current = None
    for raw in parsed_rows[1:]:
        if current is not None and len(current) < expected:
            continuation = raw[0] if raw else ''
            current[-1] = ' '.join([p for p in [current[-1], continuation] if p])
            current.extend(raw[1:])
            continue
        if current is not None:
            data_rows.append(current)
        current = list(raw)
    if current is not None:
        data_rows.append(current)
    return [headers, *data_rows]


def get_merged_numbered_column_value(base_column: str, header_indexes: dict, values: list[str]) -> str:
    base = normalize_header(base_column)
    escaped = re.escape(base)
    pattern = re.compile(rf'^{escaped}\s+\d+$')
    base_canon = canonical_header(base_column)
    canon_pattern = re.compile(rf'^{re.escape(base_canon)}\d+$')
    matched_map: dict[int, str] = {}
    for header, idx in header_indexes.items():
        if header == base or pattern.match(header):
            matched_map[idx] = header
        elif canonical_header(header) == base_canon or canon_pattern.match(canonical_header(header)):
            matched_map[idx] = header
    matched = sorted(matched_map.items(), key=lambda hi: _numeric_key(hi[1]))
    seen = []
    for idx, _ in matched:
        val = (values[idx] if idx < len(values) else '') or ''
        val = val.strip()
        if val and val not in seen:
            seen.append(val)
    return ' / '.join(seen)


def _numeric_key(header: str):
    # localeCompare numeric: split into text/number chunks
    return [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', header)]


def parse_grid(grid: list[list[str]], mode: str) -> list[dict]:
    """Map a header+rows grid into record dicts (parseCsv body)."""
    headers = grid[0] if grid else []
    record_columns = LAND_SOURCING_RECORD_COLUMNS if mode == 'land-sourcing' else BATANGAS_NEGO_RECORD_COLUMNS
    header_indexes: dict[str, int] = {}
    canon_indexes: dict[str, int] = {}
    for index, header in enumerate(headers):
        header_indexes[normalize_header(header)] = index
        canon_indexes[canonical_header(header)] = index

    rows: list[dict] = []
    for i, values in enumerate(grid[1:]):
        row: dict = {'_rowId': f'row-{i + 1}'}
        for column in record_columns:
            source_index = header_indexes.get(normalize_header(column))
            if source_index is None:
                source_index = canon_indexes.get(canonical_header(column))

            if column == 'NEGOTIATOR NAME' and source_index is None:
                source_index = header_indexes.get(normalize_header('LIST OF NEGOTIATORS'))
                if source_index is None:
                    source_index = canon_indexes.get(canonical_header('LIST OF NEGOTIATORS'))
            elif column == 'MUNICODE' and source_index is None:
                source_index = header_indexes.get(normalize_header('MUNI CODE'))
                if source_index is None:
                    source_index = canon_indexes.get(canonical_header('MUNI CODE'))

            if column == 'LSA NAME':
                row[column] = get_merged_numbered_column_value('LSA NAME', header_indexes, values)
            elif mode == 'land-sourcing' and column == 'TEAM':
                team_value = '' if source_index is None else (values[source_index] if source_index < len(values) else '')
                val = (team_value.strip()
                       or get_merged_numbered_column_value('TEAM NO.', header_indexes, values)
                       or get_merged_numbered_column_value('TEAM NO', header_indexes, values))
                row[column] = normalize_team(val)
            elif column == 'TEAM':
                val = '' if source_index is None else (values[source_index] if source_index < len(values) else '')
                row[column] = normalize_team(val)
            elif column == 'GROUP':
                val = '' if source_index is None else (values[source_index] if source_index < len(values) else '')
                row[column] = normalize_group(val)
            else:
                row[column] = '' if source_index is None else (values[source_index] if source_index < len(values) else '')
        rows.append(normalize_csv_date_columns(row))
    return rows


def validate_upload_columns(grid: list[list[str]], mode: str) -> None:
    uploaded_headers = grid[0] if grid else []
    uploaded_set = {normalize_header(h) for h in uploaded_headers}
    uploaded_canon_set = {canonical_header(h) for h in uploaded_headers}
    expected = get_expected_csv_upload_columns(mode)
    missing = []
    for column in expected:
        norm = normalize_header(column)
        canon = canonical_header(column)
        if norm in uploaded_set or canon in uploaded_canon_set:
            continue
        if column == 'NEGOTIATOR NAME':
            if normalize_header('LIST OF NEGOTIATORS') in uploaded_set or canonical_header('LIST OF NEGOTIATORS') in uploaded_canon_set:
                continue
        elif column == 'MUNICODE':
            if normalize_header('MUNI CODE') in uploaded_set or canonical_header('MUNI CODE') in uploaded_canon_set:
                continue
        missing.append(column)
    if missing:
        label = 'Land Sourcing' if mode == 'land-sourcing' else 'Negotiation'
        raise ValueError(f"CSV columns do not match the {label} template. Missing required columns: {', '.join(missing)}.")


def add_missing_upload_columns(grid: list[list[str]], mode: str) -> list[str]:
    """Append a blank column for every required upload column the file is missing.

    Lets an upload proceed (instead of hard-failing) when a required column is
    absent, and returns the canonical names of the columns that were added so the
    caller can confirm the addition to the user. An empty list means the file
    already had them all. The grid is mutated in place: the header row gains the
    new column names and every data row gains a blank cell at the same position,
    so parse_grid treats the new columns as empty values.
    """
    if not grid:
        return []
    headers = grid[0]
    uploaded_set = {normalize_header(h) for h in headers}
    uploaded_canon_set = {canonical_header(h) for h in headers}
    # Make sure every data row is at least as wide as the header before we insert.
    width = len(headers)
    for row in grid[1:]:
        if len(row) < width:
            row.extend([''] * (width - len(row)))
    added: list[str] = []
    insert_at = width
    for column in get_expected_csv_upload_columns(mode):
        norm = normalize_header(column)
        canon = canonical_header(column)
        if norm in uploaded_set or canon in uploaded_canon_set:
            continue
        if column == 'NEGOTIATOR NAME' and (
            normalize_header('LIST OF NEGOTIATORS') in uploaded_set
            or canonical_header('LIST OF NEGOTIATORS') in uploaded_canon_set
        ):
            continue
        if column == 'MUNICODE' and (
            normalize_header('MUNI CODE') in uploaded_set
            or canonical_header('MUNI CODE') in uploaded_canon_set
        ):
            continue
        added.append(column)
        headers.insert(insert_at, column)
        uploaded_set.add(norm)
        uploaded_canon_set.add(canon)
        for row in grid[1:]:
            row.insert(insert_at, '')
        insert_at += 1
    return added


# ---------------------------------------------------------------------------
# Derivations (ported one-for-one)
# ---------------------------------------------------------------------------

def get_area_index(row: dict) -> str:
    return str(row.get('AREA INDEX', row.get('AREA-INDEX', '')) or '')


def split_negotiators(value: str) -> list[str]:
    return [n.strip() for n in (value or '').split('/') if n.strip()]


def get_lo_occurrence_name(row: dict) -> str:
    return (str(row.get('REGISTERED OWNER', '') or row.get('registeredOwner', '') or '').strip()
            or str(row.get('AUTHORIZED REPRESENTATIVE', '') or row.get('authorizedRepresentative', '') or '').strip()
            or str(row.get('LO', '') or row.get('lo', '') or row.get('loName', '') or '').strip())


def _js_date_from_mm_dd_yyyy(value: str):
    """Parse an MM/DD/YYYY string the way the app relies on new Date()."""
    if not value:
        return None
    for fmt in ('%m/%d/%Y', '%m/%d/%y', '%Y-%m-%d'):
        try:
            return datetime.strptime(value.strip(), fmt)
        except ValueError:
            continue
    return None


def get_work_week_info(d: datetime) -> tuple[str, datetime, datetime]:
    """Calculate Wednesday-to-Tuesday work week (WW-01 to WW-52/53) and its start/end dates.

    The Wednesday-to-Tuesday cycle that contains Jan 1 of a year is WW-01 for that year.
    For 2026, WW-01 started on 12/31/2025 (Wed) and ran to 01/06/2026 (Tue).
    """
    w = d.weekday()  # 0=Mon, 1=Tue, 2=Wed, 3=Thu, 4=Fri, 5=Sat, 6=Sun
    days_since_wed = (w - 2) if w >= 2 else (w + 5)
    wed = d - timedelta(days=days_since_wed)
    tue = wed + timedelta(days=6)

    mid = wed + timedelta(days=3)  # Saturday midpoint determines the primary year
    y = mid.year

    jan1 = datetime(y, 1, 1)
    w_jan1 = jan1.weekday()
    days_to_wed = (w_jan1 - 2) if w_jan1 >= 2 else (w_jan1 + 5)
    ww1_wed = jan1 - timedelta(days=days_to_wed)

    week_num = int((wed - ww1_wed).days // 7) + 1
    if week_num < 1:
        y_prev = y - 1
        jan1_prev = datetime(y_prev, 1, 1)
        w_prev = jan1_prev.weekday()
        days_to_wed_prev = (w_prev - 2) if w_prev >= 2 else (w_prev + 5)
        ww1_wed_prev = jan1_prev - timedelta(days=days_to_wed_prev)
        week_num = int((wed - ww1_wed_prev).days // 7) + 1

    return f"WW-{week_num:02d}", wed, tue


def get_work_week(date_value: str) -> str:
    d = _js_date_from_mm_dd_yyyy(date_value)
    if d is None:
        return 'Unassigned'
    ww, _, _ = get_work_week_info(d)
    return ww


def work_weeks_in_range(start_str, end_str) -> list[str]:
    """The distinct Wednesday-to-Tuesday work weeks (WW-NN) that the date range
    covers, inclusive, newest first. Accepts YYYY-MM-DD or MM/DD/YYYY."""
    def _pd(s):
        s = str(s or '').strip()
        m = re.match(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$', s)
        if m:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        m = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$', s)
        if m:
            return datetime(int(m.group(3)), int(m.group(1)), int(m.group(2)))
        return None

    d0, d1 = _pd(start_str), _pd(end_str)
    if not d0 or not d1:
        return []
    if d0 > d1:
        d0, d1 = d1, d0
    labels, seen = [], set()
    _, wed, _ = get_work_week_info(d0)  # start at the Wednesday of d0's work week
    cur = wed
    while cur <= d1:
        label, _, _ = get_work_week_info(cur)
        if label not in seen:
            seen.add(label)
            labels.append(label)
        cur = cur + timedelta(days=7)
    labels.reverse()  # newest week first
    return labels


def normalize_work_week(val) -> int | None:
    """Parse a Work Week value ('WW-36', 'WW36', 'ww 36', '36', 36) into its week
    number 1..53, or None if it can't be read."""
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    m = re.search(r'(\d{1,2})', s)
    if not m:
        return None
    n = int(m.group(1))
    return n if 1 <= n <= 53 else None


def iso_work_week_num(date_value: str) -> int | None:
    """Work-week number (1..53) for a MM/DD/YYYY date using the Wed-Tue corporate work week."""
    if not date_value:
        return None
    ww_str = get_work_week(date_value)
    return normalize_work_week(ww_str)


def get_record_date_for_mode(row: dict, mode: str) -> str:
    """The date column that drives the work-week check for this workspace mode."""
    if mode == 'land-sourcing':
        return str(row.get('REPORT DATE', row.get('SOURCING DATE', '') or '') or '').strip()
    return str(row.get('NEGO DATE', '') or '').strip()


def _build_lookup_by_area(build_rows: list[dict]) -> dict:
    """area-index (lowercased) -> first build row for it (mirrors get_derived_build's first-match)."""
    lut: dict[str, dict] = {}
    for b in build_rows or []:
        area = (b.get('areaIndex', '') or '').strip().lower()
        if area and area not in lut:
            lut[area] = b
    return lut


def get_build_week_errors(rows: list[dict], build_rows: list[dict], mode: str = 'negotiation', ignore_work_week: bool = False) -> list[dict]:
    """Rows whose Build lookup is missing, or whose Build Work Week is not the same
    week or exactly one week behind the record date's work week.

    Assumes the current-year work-week numbering (get_work_week). Rows with an
    unresolved/blank date are skipped here (invalid dates are caught separately).
    When ignore_work_week is True (user opts to use current build as reference), returns [].
    Returns [{row, area, reason}].
    """
    if ignore_work_week:
        return []
    lut = _build_lookup_by_area(build_rows)
    errors: list[dict] = []
    for index, row in enumerate(rows):
        area = get_area_index(row).strip()
        if not area:
            continue
        date_val = get_record_date_for_mode(row, mode)
        date_ww = iso_work_week_num(date_val) if date_val else None
        if date_ww is None:
            continue
        b = lut.get(area.lower())
        if not b or not str(b.get('build', '') or '').strip():
            continue  # no Build for this area index is OK (defaults to BALANCE) — do not block
        prev_ww = date_ww - 1 if date_ww > 1 else date_ww
        need = f"WW-{date_ww:02d} or WW-{prev_ww:02d}"
        build_ww = normalize_work_week(b.get('workWeek'))
        if build_ww is None:
            errors.append({'row': index + 1, 'area': area,
                           'reason': f"{area}: Build has no Work Week (record is WW-{date_ww:02d}, so its Build should be {need})"})
        elif build_ww not in (date_ww, date_ww - 1):
            errors.append({'row': index + 1, 'area': area,
                           'reason': f"{area}: record is WW-{date_ww:02d} but its Build is WW-{build_ww:02d} (Build should be {need})"})
    return errors


def get_batch_work_week_bounds(rows: list[dict], mode: str, municipality_codes: list[dict] | None = None) -> tuple[str, str, set[str]]:
    """Determine (min_wednesday, max_tuesday) in YYYY-MM-DD and set of AREA-INDEXes from rows."""
    dates = []
    area_indexes = set()
    date_key = 'REPORT DATE' if mode == 'land-sourcing' else 'NEGO DATE'
    muni_list = municipality_codes or []
    for r in rows:
        val = r.get(date_key) or r.get('NEGO DATE') or r.get('REPORT DATE') or ''
        d = _js_date_from_mm_dd_yyyy(val)
        if d is not None:
            dates.append(d)
        ai = get_derived_area_index(r, muni_list).strip() if muni_list else get_area_index(r).strip()
        if not ai:
            ai = get_area_index(r).strip()
        if ai:
            area_indexes.add(ai)

    if not dates:
        now = datetime.now()
        _, min_wed, max_tue = get_work_week_info(now)
        return min_wed.strftime('%Y-%m-%d'), max_tue.strftime('%Y-%m-%d'), area_indexes

    min_d = min(dates)
    max_d = max(dates)
    _, min_wed, _ = get_work_week_info(min_d)
    _, _, max_tue = get_work_week_info(max_d)
    return min_wed.strftime('%Y-%m-%d'), max_tue.strftime('%Y-%m-%d'), area_indexes


def ordinal_duplicate(count: int) -> str:
    if 11 <= count % 100 <= 13:
        suffix = 'TH'
    elif count % 10 == 1:
        suffix = 'ST'
    elif count % 10 == 2:
        suffix = 'ND'
    elif count % 10 == 3:
        suffix = 'RD'
    else:
        suffix = 'TH'
    return f"{count}{suffix} DUPLICATE"



def _to_int(value, default=0) -> int:
    try:
        return int(float(str(value).replace(',', '')))
    except (ValueError, TypeError):
        return default


def _get_row_order_key(r: dict) -> int:
    """Return numeric row identifier or file row number from ID or _rowId."""
    rid = _to_int(r.get('ID', r.get('iD1', '0')))
    if rid > 0:
        return rid
    raw_row_id = str(r.get('_rowId') or '').strip()
    m = re.search(r'\d+', raw_row_id)
    return int(m.group()) if m else 0


def _date_sort_key(val: str) -> str:
    norm = format_date_to_mm_dd_yyyy(val)
    if norm and len(norm) == 10 and norm[2] == '/' and norm[5] == '/':
        return f"{norm[6:10]}-{norm[0:2]}-{norm[3:5]}"
    return str(val or '').strip()


def get_checker(row: dict, source_rows: list[dict], existing_records: list[dict] | None = None) -> str:
    row_id = _get_row_order_key(row)
    area_index = get_area_index(row).strip().lower()
    record_date_str = get_normalized_record_date(row)
    record_date_key = _date_sort_key(record_date_str)

    # 1. Collect all rows for this Area Index in the current uploaded batch
    all_current_area = [it for it in source_rows if get_area_index(it).strip().lower() == area_index]

    all_existing_area = []
    if existing_records:
        for it in existing_records:
            ai = (it.get('aREAINDEX') or it.get('AREA-INDEX') or it.get('AREA INDEX') or '').strip().lower()
            if ai == area_index:
                all_existing_area.append(it)

    all_area_records = all_existing_area + all_current_area

    # 2. Check same-day entries for this Area Index STRICTLY within the current uploaded batch
    if record_date_key:
        same_day_current = [it for it in all_current_area
                            if _date_sort_key(get_normalized_record_date(it)) == record_date_key]

        # Check if there is any LATER record on the exact same date in this uploaded batch
        subsequent_same_day = [it for it in same_day_current if _get_row_order_key(it) > row_id]

        if subsequent_same_day:
            # Earlier entries in the batch on the exact SAME date are duplicates (highest row# is retained)
            prior_same_day = [it for it in same_day_current if _get_row_order_key(it) < row_id]
            duplicate_num = len(prior_same_day) + 1
            return ordinal_duplicate(duplicate_num)

    # 3. This row is the first entry of the day for this Area Index.
    # Check if there was ANY earlier visit on a prior date in history
    prior_dates_records = [it for it in all_area_records
                           if record_date_key and _date_sort_key(get_normalized_record_date(it)) and _date_sort_key(get_normalized_record_date(it)) < record_date_key]

    if not prior_dates_records:
        # Truly the very first contact EVER in history -> FIRST CONTACT
        return 'FIRST CONTACT'

    # Sort prior date records by date descending, then ID descending to find the immediately preceding visit
    sorted_priors = sorted(
        prior_dates_records,
        key=lambda it: (
            _date_sort_key(get_normalized_record_date(it)),
            _to_int(it.get('ID', it.get('iD1', '0')))
        ),
        reverse=True
    )

    preceding = sorted_priors[0]
    preceding_class = str(preceding.get('classification', '') or preceding.get('CLASSIFICATION', '') or '').strip().lower()
    current_class = str(row.get('classification', '') or row.get('CLASSIFICATION', '') or '').strip().lower()

    if current_class and preceding_class and current_class != preceding_class:
        return 'REVISITED'

    return 'FOLLOW UP'


def get_year_from_row_date(row: dict) -> str:
    value = row.get('NEGO DATE') or row.get('REPORT DATE') or ''
    d = _js_date_from_mm_dd_yyyy(value)
    return '' if d is None else str(d.year)


def get_lot_area_ha(value: str) -> str:
    try:
        sqm = float(str(value or '').replace(',', ''))
    except ValueError:
        return ''
    if not math.isfinite(sqm) or sqm <= 0:
        return ''
    result = round(sqm / 10000, 4)
    # drop trailing zeros like JS Number()
    return str(result).rstrip('0').rstrip('.') if '.' in str(result) else str(result)


def get_derived_municode(row: dict, municipality_codes: list[dict]) -> str:
    raw_code = str(row.get('MUNICODE') or row.get('MuniCode') or row.get('Muni Code') or row.get('MUNI CODE') or '').strip()
    if raw_code:
        return raw_code
    muni = str(row.get('MUNICIPALITY') or row.get('Municipality') or '').strip().lower()
    prov = str(row.get('PROVINCE') or row.get('Province') or '').strip().lower()
    if not muni:
        return ''
    for code in municipality_codes:
        code_muni = str(code.get('municipality') or code.get('Municipality') or '').strip().lower()
        if code_muni and (code_muni == muni or muni.startswith(code_muni) or code_muni.startswith(muni)):
            code_prov = str(code.get('province') or code.get('Province') or '').strip().lower()
            if (not code_prov) or (not prov) or code_prov == prov or code_prov in prov or prov in code_prov:
                val = str(code.get('muniCode') or code.get('MuniCode') or code.get('Muni Code') or code.get('MUNICODE') or '').strip()
                if val:
                    return val
    return ''


def get_derived_area_index(row: dict, municipality_codes: list[dict]) -> str:
    current = get_area_index(row).strip()
    if current:
        return current
    index_no = str(row.get('INDEX NO') or row.get('Index No') or row.get('INDEX-NO') or row.get('INDEX_NO') or '').strip()
    muni_code = get_derived_municode(row, municipality_codes).strip()
    if not index_no or not muni_code:
        return ''
    return f"{muni_code}-{index_no}"


def get_normalized_record_date(row: dict) -> str:
    if not isinstance(row, dict):
        return ''
    raw = (
        row.get('NEGO DATE') or row.get('Nego Date') or row.get('nego date') or
        row.get('REPORT DATE') or row.get('Report Date') or row.get('report date') or
        row.get('SOURCING DATE') or row.get('Sourcing Date') or row.get('sourcing date') or
        row.get('SOURCED DATE') or row.get('Sourced Date') or row.get('sourced date') or
        row.get('DATE SOURCED') or row.get('Date Sourced') or row.get('date sourced') or
        row.get('DATE OF VISIT') or row.get('Date of Visit') or row.get('date of visit') or
        row.get('DATE') or row.get('Date') or row.get('date') or
        row.get('Record Date') or row.get('RECORD DATE') or
        row.get('cr63f_negotiationdate') or row.get('cr63f_reportdate') or ''
    )
    if not raw or not str(raw).strip():
        for k, v in row.items():
            if 'DATE' in str(k).upper() and v and str(v).strip():
                raw = v
                break
    if not raw or not str(raw).strip():
        return ''
    normalized = format_date_to_mm_dd_yyyy(str(raw).strip())
    return normalized.lower() if normalized else str(raw).strip().lower()


def get_derived_mapping_status(row: dict, source_rows: list[dict], existing_records: list[dict], municipality_codes: list[dict],
                               existing_area_set: set[str] | None = None,
                               batch_by_area_date: dict[tuple[str, str], list[dict]] | None = None,
                               mode: str = 'negotiation') -> str:
    area_index = (row.get('AREA INDEX') or row.get('AREA-INDEX') or get_derived_area_index(row, municipality_codes)).strip().lower()
    reg_owner = str(row.get('REGISTERED OWNER') or row.get('Registered Owner') or row.get('cr63f_registeredowner') or '').strip()
    auth_rep = str(row.get('AUTHORIZED REPRESENTATIVE') or row.get('Authorized Representative') or row.get('cr63f_authorizedrepresentative') or '').strip()

    # 1. Incomplete information check: missing area index, or missing both registered owner and authorized representative
    if not area_index or (not reg_owner and not auth_rep):
        return 'INCOMPLETE INFORMATION'

    row_date = get_normalized_record_date(row)
    row_order = _get_row_order_key(row)

    # 2. Duplicate check strictly within the uploaded file (intra-batch only)
    if row_date:
        if batch_by_area_date is not None:
            same_day_batch = batch_by_area_date.get((area_index, row_date), [])
        else:
            same_area_batch = [it for it in source_rows if (it.get('AREA INDEX') or it.get('AREA-INDEX') or get_derived_area_index(it, municipality_codes)).strip().lower() == area_index]
            same_day_batch = [it for it in same_area_batch if get_normalized_record_date(it) == row_date]
        if len(same_day_batch) > 1:
            highest_in_batch = max(same_day_batch, key=_get_row_order_key)
            if row is not highest_in_batch and _get_row_order_key(highest_in_batch) > row_order:
                return 'DUPLICATE'

    # 3. Historical check: If present in Dataverse history or existing records, it's an indexed property update
    if existing_area_set is not None:
        if area_index in existing_area_set:
            return 'INDEXED PROPERTY UPDATE'
    elif existing_records:
        for rec in existing_records:
            rec_ai = str(rec.get('aREAINDEX') or rec.get('AREA-INDEX') or rec.get('AREA INDEX') or rec.get('areaIndex') or '').strip().lower()
            if rec_ai == area_index:
                return 'INDEXED PROPERTY UPDATE'

    # 4. Otherwise brand new parcel. NEWLY MAPPED was removed as a mapping-status
    # condition, so brand-new negotiation parcels now derive NEWLY SOURCED too.
    return 'NEWLY SOURCED'


def get_derived_data_usability(row: dict, mapping_statuses: list[dict], source_rows, existing_records: list[dict], municipality_codes: list[dict], existing_area_set: set[str] | None = None, mode: str = 'negotiation') -> str:
    if source_rows is not None:
        derived = get_derived_mapping_status(row, source_rows, existing_records, municipality_codes, existing_area_set, mode=mode)
    else:
        derived = str(row.get('MAPPING STATUS') or row.get('Mapping Status') or '').strip()
    mapping_status = (derived or '').strip().lower()
    if mapping_status:
        for status in mapping_statuses:
            desc = str(status.get('description') or status.get('Description') or '').strip().lower()
            label = str(status.get('mappingLabel') or status.get('Mapping Status') or status.get('mappingStatus') or '').strip().lower()
            if (desc and desc == mapping_status) or (label and label == mapping_status):
                usability = str(status.get('dataUsability') or status.get('Data Usability') or '').strip()
                if usability:
                    return MAPPING_STATUS_DATA_USABILITY_LABEL.get(usability, usability)
    if 'DUPLICATE' in mapping_status.upper() or 'INCOMPLETE' in mapping_status.upper() or 'INCORRECT' in mapping_status.upper():
        return 'INVALID'
    if mapping_status.upper() in ('NEWLY MAPPED', 'INDEXED PROPERTY UPDATE', 'NEWLY SOURCED', 'VALID'):
        return 'VALID'
    return 'VALID' if derived else ''


# Build values that mean "this area is no longer part of the build" — treated
# exactly like an absent build row, i.e. the derived BUILD reads as BALANCE.
_BUILD_BALANCE_SENTINELS = {'NOT IN THE CURRENT BUILD', 'NOT IN CURRENT BUILD'}


def _build_value_or_balance(build_val: str) -> str:
    """A build cell's value, mapping the out-of-build sentinels to BALANCE."""
    b = (build_val or '').strip()
    if not b:
        return 'BALANCE'
    if re.sub(r'\s+', ' ', b).upper() in _BUILD_BALANCE_SENTINELS:
        return 'BALANCE'
    return b


def _record_work_week(row: dict) -> int | None:
    """ISO work-week number for a review/transaction row from its record date
    (matches the ISO numbering used for Build Work Weeks). None if undeterminable."""
    record_date = (row.get('NEGO DATE') or row.get('SOURCING DATE')
                   or row.get('REPORT DATE') or '')
    return iso_work_week_num(record_date) if record_date else None


def get_derived_build(row: dict, build_rows: list[dict], ignore_work_week: bool = False,
                      build_first_by_area: dict[str, str] | None = None,
                      build_by_area_and_ww: dict[tuple[str, int], str] | None = None) -> str:
    """Derived BUILD value for a row.

    An area-index only counts as "in the build" when it has a Build row for the
    work week IMMEDIATELY BEFORE the record's work week (WW N-1). So a WW-35
    record needs its Build at WW-34. If the area-index is absent from the Build
    table, or present only in other work weeks (e.g. only WW-33), it reads as
    BALANCE. A matched row whose value is 'NOT IN THE CURRENT BUILD' is BALANCE.

    When ignore_work_week is True (user opts to use current build as reference),
    the Area Index is matched against whatever is in build_rows without work week filtering.
    """
    area_index = get_area_index(row).strip().lower()
    if not area_index:
        return 'BALANCE'

    if build_first_by_area is not None:
        if ignore_work_week:
            return build_first_by_area.get(area_index, 'BALANCE')
        date_ww = _record_work_week(row)
        if date_ww is None:
            return build_first_by_area.get(area_index, 'BALANCE')
        accepted = (52, 53) if date_ww == 1 else (date_ww - 1,)
        for acc_ww in accepted:
            val = build_by_area_and_ww.get((area_index, acc_ww)) if build_by_area_and_ww else None
            if val:
                return val
        return 'BALANCE'

    if ignore_work_week:
        for item in build_rows:
            if (item.get('areaIndex', '') or '').strip().lower() == area_index:
                b = (item.get('build', '') or '').strip()
                if b:
                    return _build_value_or_balance(b)
        return 'BALANCE'

    date_ww = _record_work_week(row)
    if date_ww is None:
        # No usable record date -> can't place the row against a work week. Fall
        # back to a plain area-index match so the build isn't wrongly zeroed
        # while a date is still being fixed (the bad date is flagged separately).
        for item in build_rows:
            if (item.get('areaIndex', '') or '').strip().lower() == area_index:
                b = (item.get('build', '') or '').strip()
                if b:
                    return _build_value_or_balance(b)
        return 'BALANCE'

    # Accept ONLY the immediately preceding work week (WW N-1). Wrap the year
    # boundary: the week before WW-01 is the prior year's WW-52 or WW-53.
    accepted = {52, 53} if date_ww == 1 else {date_ww - 1}
    for item in build_rows:
        if (item.get('areaIndex', '') or '').strip().lower() != area_index:
            continue
        if normalize_work_week(item.get('workWeek')) not in accepted:
            continue
        b = (item.get('build', '') or '').strip()
        if b:
            return _build_value_or_balance(b)
    return 'BALANCE'


def get_nego_distinction(row: dict, rows: list[dict], existing_records: list[dict] | None = None, mode: str = 'negotiation') -> str:
    checker = str(row.get('CHECKER', '') or '')
    if 'DUPLICATE' in checker:
        return 'DUPLICATE'
    area_index = get_area_index(row).strip().lower()
    work_week = str(row.get('WORK WEEK', '') or '')
    record_date = get_normalized_record_date(row)

    same_week_current = [it for it in rows if get_area_index(it).strip().lower() == area_index and str(it.get('WORK WEEK', '') or '') == work_week]
    same_week_existing = []
    if existing_records:
        for it in existing_records:
            ai = (it.get('aREAINDEX') or it.get('AREA-INDEX') or it.get('AREA INDEX') or '').strip().lower()
            d = get_normalized_record_date(it)
            ww = str(it.get('WORK WEEK', '') or '') or (get_work_week(d) if d else '')
            if ai == area_index and ww == work_week:
                same_week_existing.append(it)

    all_week_area = same_week_existing + same_week_current
    row_date_key = _date_sort_key(record_date)
    same_date_area = [it for it in all_week_area if row_date_key and _date_sort_key(get_normalized_record_date(it)) == row_date_key]
    valid_dates = [get_normalized_record_date(it) for it in all_week_area if get_normalized_record_date(it)]
    min_date = min(valid_dates, key=_date_sort_key) if valid_dates else ''
    min_date_key = _date_sort_key(min_date)

    is_sourcing = (mode == 'land-sourcing')
    renego_single = 'REVISIT' if is_sourcing else 'RENEGO'
    renego_multi = 'REVISIT MULTIPLE' if is_sourcing else 'RENEGO MULTIPLE'

    if len(same_date_area) == 1:
        if row_date_key == min_date_key:
            return 'DISTINCT'
        renego = any(_date_sort_key(get_normalized_record_date(it)) < row_date_key and it.get('POINTS', 1) == row.get('POINTS', 1)
                     for it in all_week_area if get_normalized_record_date(it))
        return renego_single if renego else renego_multi
    return 'DISTINCT MULTIPLE' if row_date_key == min_date_key else renego_multi


def build_productivity_rows(rows: list[dict], teams: list[dict], mapping_statuses: list[dict],
                            existing_records: list[dict], municipality_codes: list[dict],
                            mode: str = 'negotiation') -> list[dict]:
    # 1. Pre-index existing records and municipality lookups for O(1) speed
    existing_by_area: dict[str, list[dict]] = collections.defaultdict(list)
    for r in (existing_records or []):
        ai = (r.get('aREAINDEX') or r.get('AREA-INDEX') or r.get('AREA INDEX') or '').strip().lower()
        if ai:
            existing_by_area[ai].append(r)

    # Pre-index existing Dataverse records by Work Week + Landowner for LO WW occurrences
    batch_ids = {str(r.get('ID', '') or r.get('iD1', '')).strip() for r in rows if (r.get('ID') or r.get('iD1'))}
    existing_lo_ww_counts = collections.Counter()
    seen_existing_keys = set()
    for r in (existing_records or []):
        r_id = str(r.get('ID', '') or r.get('iD1', '') or r.get('_record_id', '')).strip()
        r_ai = (r.get('aREAINDEX') or r.get('AREA-INDEX') or r.get('AREA INDEX') or '').strip().lower()
        r_date = r.get('NEGO DATE') or r.get('REPORT DATE') or r.get('date') or ''
        r_key = r_id or f"{r_ai}_{r_date}"
        if r_key in seen_existing_keys:
            continue
        seen_existing_keys.add(r_key)

        # Skip if this record is already in the current batch being processed
        if r_id and r_id in batch_ids:
            continue

        r_lo = get_lo_occurrence_name(r)
        r_lo_key = r_lo.strip().lower()
        r_ww = r.get('workWeek') or r.get('WORK WEEK') or get_work_week(r_date)
        if r_ww and r_ww != 'Unassigned' and r_lo_key:
            existing_lo_ww_counts[(r_ww, r_lo_key)] += 1

    teams_by_lower = {(item.get('employeeName', '') or '').strip().lower(): item for item in teams if item.get('employeeName')}

    source_with_area = [{**sr, 'AREA INDEX': get_derived_area_index(sr, municipality_codes),
                         'AREA-INDEX': get_derived_area_index(sr, municipality_codes)} for sr in rows]

    source_by_area: dict[str, list[dict]] = collections.defaultdict(list)
    for sr in source_with_area:
        ai = str(sr.get('AREA INDEX') or '').strip().lower()
        if ai:
            source_by_area[ai].append(sr)

    expanded: list[dict] = []
    for row in rows:
        derived_area_index = get_derived_area_index(row, municipality_codes)
        ai_key = derived_area_index.strip().lower()
        record_date = row.get('NEGO DATE') or row.get('REPORT DATE') or ''
        negotiators = split_negotiators(row.get('NEGOTIATOR NAME') or row.get('LSA NAME') or '')
        names_to_render = negotiators if negotiators else ['']
        checker_row = {**row, 'AREA INDEX': derived_area_index, 'AREA-INDEX': derived_area_index}
        
        # Scoped to matching Area Index for O(1) lookup
        checker = get_checker(checker_row, source_by_area.get(ai_key, []), existing_by_area.get(ai_key, []))
        data_usability = row.get('DATA USABILITY') or get_derived_data_usability(row, mapping_statuses, None, existing_records, municipality_codes)

        for name in names_to_render:
            team = teams_by_lower.get((name or '').strip().lower()) if name else None
            negotiator_match = 'MATCH' if team else 'NO MATCH'
            corr_team = normalize_team(team.get('team', '')) if team else ('Unmatched' if name else 'TEAM 0')
            corr_group = normalize_group(team.get('group', '')) if team else ('Unmatched' if name else 'GROUP 0')
            expanded.append({
                'NEGO DATE': record_date, 'AREA INDEX': derived_area_index, 'PROVINCE': row.get('PROVINCE', '') or '',
                'MUNICIPALITY': row.get('MUNICIPALITY', '') or '', 'BARANGAY': row.get('BARANGAY', '') or '',
                'TYPE OF REPORT': row.get('TYPE OF REPORT', '') or '',
                'ACTION': row.get('ACTION', '') or '', 'WORK WEEK': get_work_week(record_date),
                'TEAM': normalize_team(row.get('TEAM', '')),
                'GROUP': normalize_group(row.get('GROUP', '')),
                'NEGOTIATOR NAME': name, 'NEGO CODE': row.get('NEGO CODE', '') or '',
                'NEGOTIATOR MATCH': negotiator_match, 'DATA USABILITY': data_usability,
                'UNIQUE ID': f"{derived_area_index}{record_date}{checker}", 'CHECKER': checker, 'POINTS': 0,
                'NEGO DISTINCTION': '', 'LO': get_lo_occurrence_name(row), 'LO OCCURRENCE KEY': get_lo_occurrence_name(row),
                'LO OCCURRENCE BY DAY': 0, 'LO OCCURRENCE BY WW': 0, 'LO POINTS BY DAY': 0, 'LO POINTS BY WW': 0,
                'LO COUNT BY DAY': 0, 'LO COUNT BY WW': 0,
                'CORRECT TEAM': corr_team,
                'CORRECT GROUP': corr_group,
            })

    # 2. O(N) pre-aggregation for Points and LO counts
    uid_counts = collections.Counter(r.get('UNIQUE ID') for r in expanded)
    rows_with_points = [{**r, 'POINTS': to_fixed(1 / (uid_counts.get(r.get('UNIQUE ID'), 1) or 1), 4)} for r in expanded]

    # Pre-index rows_with_points and existing_records by (area_index, work_week) for instant distinction lookups
    rwp_by_area_ww: dict[tuple[str, str], list[dict]] = collections.defaultdict(list)
    for r in rows_with_points:
        ai = str(r.get('AREA INDEX') or '').strip().lower()
        ww = str(r.get('WORK WEEK', '') or '')
        if ai:
            rwp_by_area_ww[(ai, ww)].append(r)

    existing_by_area_ww: dict[tuple[str, str], list[dict]] = collections.defaultdict(list)
    for r in (existing_records or []):
        ai = (r.get('aREAINDEX') or r.get('AREA-INDEX') or r.get('AREA INDEX') or '').strip().lower()
        d = get_normalized_record_date(r)
        ww = str(r.get('WORK WEEK', '') or '') or (get_work_week(d) if d else '')
        if ai:
            existing_by_area_ww[(ai, ww)].append(r)

    batch_lo_day_counts = collections.Counter(
        (r.get('NEGO DATE') or r.get('REPORT DATE') or '', get_lo_occurrence_name(r).strip().lower())
        for r in rows if get_lo_occurrence_name(r)
    )
    batch_lo_ww_counts = collections.Counter(
        (get_work_week(r.get('NEGO DATE') or r.get('REPORT DATE') or ''), get_lo_occurrence_name(r).strip().lower())
        for r in rows if get_lo_occurrence_name(r)
    )

    result = []
    for row in rows_with_points:
        lo_name = str(row.get('LO OCCURRENCE KEY', '') or '')
        lo_key = lo_name.strip().lower()
        d_val = row.get('NEGO DATE')
        ww_val = str(row.get('WORK WEEK', '') or '')
        ai_key = str(row.get('AREA INDEX') or '').strip().lower()

        day_count = batch_lo_day_counts.get((d_val, lo_key), 1) if lo_key else 1
        batch_week_count = batch_lo_ww_counts.get((ww_val, lo_key), 1) if lo_key else 1
        existing_week_count = existing_lo_ww_counts.get((ww_val, lo_key), 0) if lo_key else 0
        week_count = batch_week_count + existing_week_count
        if week_count < 1:
            week_count = 1

        lo_points_by_day = to_fixed(1 / day_count, 4)
        lo_points_by_ww = to_fixed(1 / week_count, 4)

        result.append({
            **row,
            'NEGO DISTINCTION': get_nego_distinction(row, rwp_by_area_ww.get((ai_key, ww_val), []), existing_by_area_ww.get((ai_key, ww_val), []), mode),
            'LO OCCURRENCE BY DAY': day_count, 'LO OCCURRENCE BY WW': week_count,
            'LO POINTS BY DAY': lo_points_by_day, 'LO POINTS BY WW': lo_points_by_ww,
            'LO COUNT BY DAY': to_fixed(day_count * lo_points_by_day, 4), 'LO COUNT BY WW': 1,
        })
    return result


# ---------------------------------------------------------------------------
# Workspace enrichment (matched negotiator, team/group match, item #)
# ---------------------------------------------------------------------------

def normalize_name_part(value: str) -> str:
    return re.sub(r'[^a-z]', '', (value or '').lower())


def get_fuzzy_negotiator_match(source_name: str, options: list[str]) -> str:
    trimmed = (source_name or '').strip()
    if not trimmed:
        return ''
    compact = normalize_name_part(trimmed)
    for option in options:
        if normalize_name_part(option) == compact:
            return option
    source_parts = [normalize_name_part(p) for p in re.split(r'[\s._,-]+', trimmed.lower()) if normalize_name_part(p)]
    if len(source_parts) < 2:
        return ''
    first_initial = source_parts[0][0]
    last_token = source_parts[-1]
    for option in options:
        option_parts = [normalize_name_part(p) for p in re.split(r'[\s._,-]+', option.lower()) if normalize_name_part(p)]
        if not option_parts:
            continue
        option_first = option_parts[0]
        option_last = option_parts[-1]
        if option_first and option_last and option_first.startswith(first_initial) and option_last == last_token:
            return option
    return ''


def get_match_status(source_value, lookup_value) -> str:
    s = str(source_value if source_value is not None else '').strip().lower()
    l = str(lookup_value if lookup_value is not None else '').strip().lower()
    if not s and not l:
        return ''
    return 'MATCH' if s == l else 'NO MATCH'


def enrich_workspace_rows(productivity_rows: list[dict], teams: list[dict]) -> list[dict]:
    negotiator_options = sorted({(t.get('employeeName', '') or '').strip() for t in teams if (t.get('employeeName', '') or '').strip()})
    lookup = build_team_lookup(teams)

    enriched = []
    for index, row in enumerate(productivity_rows):
        matched = str(row.get('MATCHED NEGOTIATOR NAME', '') or '').strip()
        if not matched:
            matched = get_fuzzy_negotiator_match(str(row.get('NEGOTIATOR NAME', '') or ''), negotiator_options)
        defaults = lookup.get(matched.strip().lower(), {'correctTeam': 'TEAM 0', 'correctGroup': 'GROUP 0'})
        team_val = normalize_team(row.get('TEAM'))
        group_val = normalize_group(row.get('GROUP'))
        enriched.append({
            **row,
            'ITEM #': row.get('ITEM #', index + 1),
            'MATCHED NEGOTIATOR NAME': matched,
            'TEAM': team_val,
            'CORRECT TEAM': defaults['correctTeam'],
            'TEAM MATCH': get_match_status(team_val, defaults['correctTeam']),
            'GROUP': group_val,
            'CORRECT GROUP': defaults['correctGroup'],
            'GROUP MATCH': get_match_status(group_val, defaults['correctGroup']),
            # workspace loadProductivityRows rounds LO COUNT BY DAY to a whole number
            'LO COUNT BY DAY': round_whole_number_text(row.get('LO COUNT BY DAY')),
        })
    return enriched


# ---------------------------------------------------------------------------
# Output column layout
# ---------------------------------------------------------------------------

OUTPUT_COLUMNS = [
    'ITEM #', 'AREA INDEX', 'NEGO DATE', 'WORK WEEK', 'PROVINCE', 'MUNICIPALITY', 'BARANGAY', 'TYPE OF REPORT', 'ACTION',
    'NEGOTIATOR NAME', 'NEGO CODE', 'MATCHED NEGOTIATOR NAME', 'TEAM', 'CORRECT TEAM', 'TEAM MATCH', 'GROUP',
    'CORRECT GROUP', 'GROUP MATCH', 'DATA USABILITY', 'UNIQUE ID', 'CHECKER', 'NEGO DISTINCTION', 'LO', 'POINTS',
    'LO OCCURRENCE BY DAY', 'LO OCCURRENCE BY WW', 'LO POINTS BY DAY', 'LO POINTS BY WW', 'LO COUNT BY DAY', 'LO COUNT BY WW',
]

LAND_SOURCING_HIDDEN = {'NEGO CODE', 'ACTION'}
NEGOTIATION_HIDDEN = {'NEGO CODE'}  # NEGO CODE kept? app hides only in visibleProductivityColumns; keep for export


def get_output_columns(mode: str) -> list[str]:
    hidden = LAND_SOURCING_HIDDEN if mode == 'land-sourcing' else set()
    return [c for c in OUTPUT_COLUMNS if c not in hidden]


def get_header_label(column: str, mode: str) -> str:
    if column == 'ITEM #':
        return '#'
    if column == 'MATCHED NEGOTIATOR NAME':
        return 'MATCHED SOURCER' if mode == 'land-sourcing' else 'MATCHED NEGOTIATOR'
    if mode == 'land-sourcing':
        if column == 'NEGO DATE':
            return 'SOURCING DATE'
        if column == 'NEGOTIATOR NAME':
            return 'SOURCER NAME'
        if column == 'NEGO DISTINCTION':
            return 'SOURCING DISTINCTION'
    return column


# ---------------------------------------------------------------------------
# File IO
# ---------------------------------------------------------------------------

def read_grid_from_file(path: str, mode: str) -> list[list[str]]:
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.csv', '.tsv', '.txt'):
        try:
            with open(path, 'r', encoding='utf-8-sig', newline='') as f:
                text = f.read()
        except UnicodeDecodeError:
            with open(path, 'r', encoding='cp1252', newline='', errors='replace') as f:
                text = f.read()
        parsed = parse_delimited_text(text)
        return coalesce_parsed_rows(parsed, mode)
    if ext in ('.xlsx', '.xlsm'):
        if load_workbook is None:
            raise RuntimeError('openpyxl is required to read Excel files. Install it with: pip install openpyxl')
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        grid = []
        for r in ws.iter_rows(values_only=True):
            grid.append(['' if c is None else str(c).strip() for c in r])
        wb.close()
        # drop fully-empty rows (matches CSV parser behaviour)
        grid = [row for row in grid if any(v for v in row)]
        return grid
    raise ValueError(f'Unsupported input format: {ext}. Use .csv or .xlsx.')


def _sheet_rows_to_dicts(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_header('' if h is None else str(h)) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(c is not None and str(c).strip() for c in r):
            continue
        d = {}
        for i, h in enumerate(headers):
            d[h] = '' if i >= len(r) or r[i] is None else str(r[i]).strip()
        out.append(d)
    return out


def _pick(d: dict, *names) -> str:
    for name in names:
        key = normalize_header(name)
        if key in d and str(d[key]).strip():
            return str(d[key]).strip()
        canon = canonical_header(name)
        for dk, dv in d.items():
            if canonical_header(dk) == canon and str(dv).strip():
                return str(dv).strip()
    return ''


def load_reference_workbook(path: str):
    """Return (teams, municipality_codes, mapping_statuses, build_rows)."""
    if load_workbook is None:
        raise RuntimeError('openpyxl is required to read the reference workbook.')
    wb = load_workbook(path, read_only=True, data_only=True)
    teams, municipality_codes, mapping_statuses, build_rows = [], [], [], []
    for ws in wb.worksheets:
        title = normalize_header(ws.title)
        dicts = _sheet_rows_to_dicts(ws)
        if 'TEAM' in title and 'COMPOSITION' in title or title.startswith('TEAM'):
            for d in dicts:
                teams.append({'employeeName': _pick(d, 'Employee Name', 'Name', 'Negotiator', 'LSA Name', 'Negotiator Name'),
                              'team': normalize_team(_pick(d, 'Team')), 'group': normalize_group(_pick(d, 'Group')),
                              'dept': _pick(d, 'Department', 'Dept', 'DEPARTMENT', 'DEPT')})
        elif 'MUNICIPAL' in title or 'MUNI' in title:
            for d in dicts:
                municipality_codes.append({'municipality': _pick(d, 'Municipality'),
                                           'muniCode': _pick(d, 'MuniCode', 'Muni Code', 'Municode', 'Code'),
                                           'province': _pick(d, 'Province')})
        elif 'MAPPING' in title:
            for d in dicts:
                mapping_statuses.append({'description': _pick(d, 'Description'),
                                         'mappingLabel': _pick(d, 'Mapping Status', 'Status', 'Label'),
                                         'dataUsability': _pick(d, 'Data Usability', 'Usability', 'DataUsability')})
        elif 'BUILD' in title:
            for d in dicts:
                build_rows.append({'areaIndex': _pick(d, 'Area Index', 'Area-Index', 'AreaIndex'),
                                   'build': _pick(d, 'Build'),
                                   'workWeek': _pick(d, 'Work Week', 'WORK WEEK', 'WorkWeek', 'WW', 'Work-Week')})
    wb.close()
    teams = [t for t in teams if t['employeeName']]
    municipality_codes = [m for m in municipality_codes if m['municipality']]
    mapping_statuses = [m for m in mapping_statuses if m['description'] or m['mappingLabel']]
    build_rows = [b for b in build_rows if b['areaIndex']]
    return teams, municipality_codes, mapping_statuses, build_rows


# Per-source reference table definitions ------------------------------------
REFERENCE_TEMPLATE_HEADERS = {
    'team': ['Employee Name', 'Team', 'Group', 'Department'],
    'municipality': ['Municipality', 'MuniCode', 'Province'],
    'mapping': ['Description', 'Mapping Status', 'Data Usability'],
    'build': ['Area Index', 'Build', 'Work Week'],
    'existing': ['ID', 'AREA-INDEX'],
}

REFERENCE_VIEW_COLUMNS = {
    'team': [('employeeName', 'Employee Name'), ('team', 'Team'), ('group', 'Group'), ('dept', 'Department')],
    'municipality': [('municipality', 'Municipality'), ('muniCode', 'MuniCode'), ('province', 'Province')],
    'mapping': [('description', 'Description'), ('mappingLabel', 'Mapping Status'), ('dataUsability', 'Data Usability')],
    'build': [('areaIndex', 'Area Index'), ('build', 'Build'), ('workWeek', 'Work Week')],
    'existing': [('iD1', 'ID'), ('aREAINDEX', 'AREA-INDEX'), ('date', 'Record Date'), ('classification', 'Classification')],
}

REFERENCE_LABELS = {
    'team': 'Team Composition', 'municipality': 'Municipality Code',
    'mapping': 'Mapping Status', 'build': 'Build Table', 'existing': 'Existing Records',
}


def _read_records_from_file(path: str) -> list[dict]:
    """Read a single-table CSV/Excel into rows keyed by normalized header."""
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.csv', '.tsv', '.txt'):
        with open(path, 'r', encoding='utf-8-sig', newline='') as f:
            grid = parse_delimited_text(f.read())
        if not grid:
            return []
        headers = [normalize_header(h) for h in grid[0]]
        out = []
        for r in grid[1:]:
            if not any((c or '').strip() for c in r):
                continue
            out.append({headers[i]: (r[i] if i < len(r) else '') for i in range(len(headers))})
        return out
    if ext in ('.xlsx', '.xlsm'):
        if load_workbook is None:
            raise RuntimeError('openpyxl is required to read Excel files.')
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = _sheet_rows_to_dicts(wb.active)
        wb.close()
        return rows
    raise ValueError(f'Unsupported reference format: {ext}. Use .csv or .xlsx.')


def load_single_reference(path: str, kind: str) -> list[dict]:
    """Load one reference table (team/municipality/mapping/build/existing) from a file."""
    rows = _read_records_from_file(path)
    out: list[dict] = []
    if kind == 'team':
        for d in rows:
            out.append({
                'employeeName': _pick(d, 'Employee Name', 'EMPLOYEE NAME', 'Employee_Name', 'Name', 'NAME', 'Employee', 'EMPLOYEE'),
                'team': normalize_team(_pick(d, 'Team', 'TEAM', 'Team Name', 'TEAM NAME')),
                'group': normalize_group(_pick(d, 'Group', 'GROUP', 'Group Name', 'GROUP NAME')),
                'dept': _pick(d, 'Department', 'DEPARTMENT', 'Dept', 'DEPT', 'Department Name'),
            })
        return [r for r in out if r['employeeName']]
    if kind == 'municipality':
        for d in rows:
            out.append({
                'municipality': _pick(d, 'Municipality', 'MUNICIPALITY', 'City', 'CITY', 'Town', 'TOWN'),
                'muniCode': _pick(d, 'MuniCode', 'Muni Code', 'MUNICODE', 'MUNI CODE', 'Muni_Code', 'Code', 'CODE'),
                'province': _pick(d, 'Province', 'PROVINCE', 'Prov', 'PROV'),
            })
        return [r for r in out if r['muniCode']]
    if kind == 'mapping':
        for d in rows:
            out.append({
                'description': _pick(d, 'Description', 'DESCRIPTION', 'Desc', 'DESC', 'Mapping Status', 'MAPPING STATUS'),
                'mappingLabel': _pick(d, 'Mapping Status', 'MAPPING STATUS', 'MappingStatus', 'Status', 'STATUS'),
                'dataUsability': _pick(d, 'Data Usability', 'DATA USABILITY', 'DataUsability', 'Usability', 'USABILITY'),
            })
        return [r for r in out if r['mappingLabel']]
    if kind == 'build':
        for d in rows:
            area = _pick(d, 'Area Index', 'AREA-INDEX', 'AREA INDEX', 'AreaIndex', 'Area-Index')
            if not area:
                muni = _pick(d, 'MuniCode', 'Muni Code', 'MUNICODE', 'MUNI CODE', 'Muni_Code', 'Municipality', 'MUNICIPALITY')
                idx = _pick(d, 'Index No', 'Index-No', 'INDEX NO', 'INDEX-NO', 'Index_No', 'INDEX_NO', 'Lot No', 'LOT NO')
                if muni and idx:
                    area = f"{muni}-{idx}"
            build_val = _pick(d, 'Build', 'BUILD', 'Build Status', 'BUILD STATUS', 'Build_Status', 'Status', 'STATUS')
            ww_val = _pick(d, 'Work Week', 'WORK WEEK', 'WorkWeek', 'WW', 'Work-Week', 'WORKWEEK')
            out.append({'areaIndex': area, 'build': build_val, 'workWeek': ww_val})
        return [r for r in out if r['areaIndex']]
    if kind == 'existing':
        for d in rows:
            area = _pick(d, 'AREA-INDEX', 'Area Index', 'AreaIndex', 'AREA INDEX')
            raw_date = _pick(d, 'Record Date', 'RECORD DATE', 'NEGO DATE', 'Nego Date', 'REPORT DATE', 'Report Date', 'SOURCING DATE', 'Sourcing Date', 'DATE', 'Date')
            classification = _pick(d, 'Classification', 'CLASSIFICATION', 'Class', 'CLASS')
            if area:
                out.append({
                    'aREAINDEX': area,
                    'iD1': _to_int(_pick(d, 'ID', 'ID1', 'ITEM #', 'ITEM NO')),
                    'date': raw_date,
                    'NEGO DATE': raw_date,
                    'REPORT DATE': raw_date,
                    'classification': classification,
                    'CLASSIFICATION': classification,
                })
        return out
    raise ValueError(f'Unknown reference kind: {kind}')


def write_reference_template(path: str, kind: str) -> None:
    headers = [label for _, label in REFERENCE_VIEW_COLUMNS.get(kind, [])]
    if not headers:
        raise ValueError(f'Unknown reference kind: {kind}')
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        if Workbook is None:
            raise RuntimeError('openpyxl is required to write Excel templates.')
        wb = Workbook()
        ws = wb.active
        ws.title = REFERENCE_LABELS.get(kind, kind)[:31]
        ws.append(headers)
        wb.save(path)
    else:
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            csv.writer(f).writerow(headers)


def load_existing_records(path: str) -> list[dict]:
    """Existing record rows -> [{aREAINDEX, iD1, date, classification}] for INDEXED detection + latest ID."""
    ext = os.path.splitext(path)[1].lower()
    records = []
    if ext in ('.csv', '.tsv', '.txt'):
        with open(path, 'r', encoding='utf-8-sig', newline='') as f:
            grid = parse_delimited_text(f.read())
        if grid:
            headers = [normalize_header(h) for h in grid[0]]
            for r in grid[1:]:
                d = {headers[i]: (r[i] if i < len(r) else '') for i in range(len(headers))}
                raw_d = _pick(d, 'Record Date', 'RECORD DATE', 'NEGO DATE', 'Nego Date', 'REPORT DATE', 'Report Date', 'SOURCING DATE', 'Sourcing Date', 'DATE', 'Date')
                cls_v = _pick(d, 'Classification', 'CLASSIFICATION', 'Class', 'CLASS')
                records.append({
                    'aREAINDEX': _pick(d, 'AREA-INDEX', 'AREA INDEX', 'AreaIndex'),
                    'iD1': _to_int(_pick(d, 'ID', 'ID1', 'ITEM #', 'ITEM NO')),
                    'date': raw_d,
                    'NEGO DATE': raw_d,
                    'REPORT DATE': raw_d,
                    'classification': cls_v,
                    'CLASSIFICATION': cls_v,
                })
    elif ext in ('.xlsx', '.xlsm'):
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for d in _sheet_rows_to_dicts(ws):
                area = _pick(d, 'AREA-INDEX', 'AREA INDEX', 'AreaIndex')
                raw_d = _pick(d, 'Record Date', 'RECORD DATE', 'NEGO DATE', 'Nego Date', 'REPORT DATE', 'Report Date', 'SOURCING DATE', 'Sourcing Date', 'DATE', 'Date')
                cls_v = _pick(d, 'Classification', 'CLASSIFICATION', 'Class', 'CLASS')
                if area:
                    records.append({
                        'aREAINDEX': area,
                        'iD1': _to_int(_pick(d, 'ID', 'ID1', 'ITEM #', 'ITEM NO')),
                        'date': raw_d,
                        'NEGO DATE': raw_d,
                        'REPORT DATE': raw_d,
                        'classification': cls_v,
                        'CLASSIFICATION': cls_v,
                    })
        wb.close()
    return records


def apply_sequential_ids(rows: list[dict], latest_id: int) -> list[dict]:
    return [{**row, 'ID': str(latest_id + i + 1)} for i, row in enumerate(rows)]


def write_output(rows: list[dict], path: str, mode: str) -> None:
    columns = get_output_columns(mode)
    labels = [get_header_label(c, mode) for c in columns]
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(labels)
            for row in rows:
                writer.writerow([js_number_str(row.get(c, '')) for c in columns])
    elif ext in ('.xlsx', '.xlsm'):
        if Workbook is None:
            raise RuntimeError('openpyxl is required to write Excel files.')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Productivity'
        ws.append(labels)
        for row in rows:
            ws.append([_cell(row.get(c, '')) for c in columns])
        wb.save(path)
    else:
        raise ValueError(f'Unsupported output format: {ext}. Use .csv or .xlsx.')


def write_review_grid(rows: list[dict], path: str, mode: str) -> None:
    """Write the editable CSV Workspace rows back into a CSV/Excel file.

    Backs the desktop app's "Save to file" action so in-app edits and the
    calculated columns are reflected straight back into the uploaded transaction
    data file. Column order and headers match the review table
    (get_review_headers) so the saved file round-trips cleanly through the app's
    upload/review flow. Only .csv and .xlsx are supported (an .xlsm path is
    written as .xlsx, matching read_grid_from_file's accepted inputs).
    """
    columns = get_review_headers(mode)
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            for row in rows:
                writer.writerow([js_number_str(row.get(c, '')) for c in columns])
    elif ext in ('.xlsx', '.xlsm'):
        if Workbook is None:
            raise RuntimeError('openpyxl is required to write Excel files.')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Data'
        ws.append(columns)
        for row in rows:
            ws.append([_cell(row.get(c, '')) for c in columns])
        wb.save(path)
    else:
        raise ValueError(f'Unsupported file format: {ext}. Use .csv or .xlsx.')


def _cell(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    return '' if value is None else str(value)


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

def get_review_headers(mode: str) -> list[str]:
    """Column order shown in the editable review table (headers useMemo)."""
    active = LAND_SOURCING_RECORD_COLUMNS if mode == 'land-sourcing' else BATANGAS_NEGO_RECORD_COLUMNS
    if mode == 'land-sourcing':
        priority = ['ID', 'AREA-INDEX', 'INDEX NO', 'REPORT DATE', 'DATE UPLOADED', 'DATE EXTRACTED', 'LSA NAME', 'TEAM', 'GROUP']
    else:
        priority = ['ID', 'AREA-INDEX', 'INDEX NO', 'NEGO DATE', 'NEGOTIATOR NAME', 'TEAM', 'GROUP']
    return priority + [c for c in active if c not in priority]


def dept_for_mode(mode: str) -> str:
    """Owning department for a load mode: Negotiation -> COMDEV, Land Sourcing -> LSR."""
    return 'LSR' if mode == 'land-sourcing' else 'COMDEV'


def filter_teams_by_mode(teams: list[dict], mode: str) -> list[dict]:
    """Restrict team-composition rows to the department that owns the given mode
    (Negotiation -> COMDEV, Land Sourcing -> LSR). Rows whose Department is blank
    are kept so a reference table that has not populated the Department column
    never empties the Team / Group / Name dropdowns."""
    target = dept_for_mode(mode)
    out = []
    for t in teams:
        dept = (t.get('dept', '') or '').strip().upper()
        if not dept or target in dept:
            out.append(t)
    return out


def get_column_type_and_choices(col: str, mode: str, teams: list[dict],
                                municipality_codes: list[dict], mapping_statuses: list[dict],
                                build_rows: list[dict]) -> dict:
    col_upper = col.upper().strip()

    # 1. Numbers (Integer)
    if col_upper in ('ID', 'ITEM #', 'YEAR', 'ADVANCE PAYMENT (YRS)', 'LO OCCURRENCE BY DAY', 'LO OCCURRENCE BY WW', 'LO COUNT BY DAY', 'LO COUNT BY WW'):
        return {'type': 'number'}

    # 2. Decimals / Floats
    if col_upper in ('LOT AREA (SQM)', 'LOT AREA (HA)', 'XCOORD', 'YCOORD', 'POINTS',
                     'LO POINTS BY DAY', 'LO POINTS BY WW',
                     'NEGO INDICATIVE PRICE - ASKING (SALE)', 'NEGO INDICATIVE PRICE - OFFER (SALE)',
                     'NEGO FIRST PAYMENT', 'NEGO INDICATIVE PRICE - ASKING (LEASE)',
                     'NEGO INDICATIVE PRICE - OFFER (LEASE)', 'INDICATIVE PRICE SALE',
                     'INDICATIVE PRICE LEASE', 'NEGO LEASE ESCALATION RATE', 'ESCALATION'):
        return {'type': 'decimal'}

    # 3. Dates
    if col_upper in ('NEGO DATE', 'REPORT DATE', 'SOURCING DATE', 'DATE UPLOADED', 'DATE EXTRACTED', 'NEXT CALL DATE'):
        return {'type': 'date'}

    # 4. Choices (Picklists)
    if col_upper in ('MAPPING STATUS', 'MAPPING-STATUS'):
        # NEWLY MAPPED was removed as a mapping-status condition. Exclude it from
        # both the Dataverse reference-table labels and the standard set.
        excluded = {'UNASSIGNED', 'NEWLY MAPPED'}
        choices = []
        for ms in mapping_statuses:
            lbl = ms.get('mappingLabel') or ms.get('mappingStatus') or ms.get('Mapping Status') or ms.get('status') or ms.get('label')
            if lbl and lbl.upper() not in excluded and lbl not in choices:
                choices.append(lbl)
        for std in ['NEWLY SOURCED', 'INDEXED PROPERTY UPDATE', 'DUPLICATE', 'INCOMPLETE INFORMATION', 'INCORRECT GEOTAG']:
            if std not in choices:
                choices.append(std)
        return {'type': 'choice', 'choices': choices}

    if col_upper in ('DATA USABILITY', 'DATA-USABILITY'):
        return {'type': 'choice', 'choices': ['VALID', 'INVALID']}

    if col_upper in ('TEAM MATCH', 'GROUP MATCH'):
        return {'type': 'choice', 'choices': ['MATCH', 'NO MATCH']}

    if col_upper in ('TEAM', 'TEAM NAME'):
        team_choices = []
        for t in filter_teams_by_mode(teams, mode):
            tv = normalize_team(t.get('team', ''))
            if tv and tv not in team_choices:
                team_choices.append(tv)
        if 'TEAM 0' not in team_choices:
            team_choices.append('TEAM 0')
        return {'type': 'choice', 'choices': sorted(team_choices)}

    if col_upper in ('GROUP', 'GROUP NAME'):
        group_choices = []
        for t in filter_teams_by_mode(teams, mode):
            gv = normalize_group(t.get('group', ''))
            if gv and gv not in group_choices:
                group_choices.append(gv)
        if 'GROUP 0' not in group_choices:
            group_choices.append('GROUP 0')
        return {'type': 'choice', 'choices': sorted(group_choices)}

    if col_upper == 'CHECKER':
        return {'type': 'choice', 'choices': ['FIRST CONTACT', 'FOLLOW UP', 'REVISITED', '1ST DUPLICATE', '2ND DUPLICATE', '3RD DUPLICATE']}

    if col_upper in ('SOURCING DISTINCTION', 'NEGO DISTINCTION', 'DISTINCTION'):
        if mode == 'land-sourcing':
            return {'type': 'choice', 'choices': ['DISTINCT', 'REVISIT', 'DISTINCT MULTIPLE', 'REVISIT MULTIPLE', 'DUPLICATE']}
        else:
            return {'type': 'choice', 'choices': ['DISTINCT', 'RENEGO', 'DISTINCT MULTIPLE', 'RENEGO MULTIPLE', 'DUPLICATE']}

    if col_upper == 'WORK WEEK':
        return {'type': 'choice', 'choices': [f'WW-{i:02d}' for i in range(1, 54)]}

    if col_upper == 'BUILD':
        choices = []
        for b in build_rows:
            bv = (b.get('build', '') or '').strip()
            if bv and bv.upper() != 'UNASSIGNED' and bv not in choices:
                choices.append(bv)
        for std in ['BUILD-1', 'BUILD-2', 'BUILD-3', 'BALANCE']:
            if std not in choices:
                choices.append(std)
        return {'type': 'choice', 'choices': choices}

    if col_upper in ('NET OR GROSS (SALE)', 'NET OR GROSS (LEASE)', 'NET OR GROSS'):
        return {'type': 'choice', 'choices': ['NET', 'GROSS']}

    if col_upper == 'NEGO CODE':
        return {'type': 'choice', 'choices': ['ASSOCIATE', 'OFFICER', 'SENIOR OFFICER', 'MANAGER', 'SENIOR MANAGER', 'COO']}

    if col_upper in ('CLASSIFICATION', 'CLASSIFICATION STATUS', 'CLASSIFICATIONSTATUS'):
        if mode == 'land-sourcing':
            return {'type': 'choice', 'choices': [
                'FOR VERIFICATION',
                'NOT AVAILABLE',
                'NOT USABLE',
                'SECURED SALE',
                'SECURED LEASE',
                'OPEN TO LEASE',
                'OPEN TO SALE',
            ]}
        else:
            return {'type': 'choice', 'choices': [
                'NEGOTIATING - SALE',
                'NEGOTIATING - LEASE',
                'CLOSED - SALE',
                'CLOSED - LEASE',
                'NOT OPEN - WILL NOT SELL',
                'NOT OPEN - ESTATE',
                'DROPPED',
                'FOR VERIFICATION',
                'NOT AVAILABLE',
                'NOT USABLE',
                'SECURED SALE',
                'SECURED LEASE',
                'OPEN TO LEASE',
                'OPEN TO SALE',
                'AWAITING SIGNED LEASE',
                'AWAITING SIGNED SALE',
            ]}

    if col_upper == 'CLASS':
        # CLASS is a dependent dropdown: the options actually offered depend on
        # the row's CLASSIFICATION value (handled in the UI). This list is the
        # union of every group so validation accepts any legitimate value.
        return {'type': 'choice', 'dependsOn': 'CLASSIFICATION', 'choices': [
            # NA (NOT AVAILABLE)
            '1-SOURCE OF INCOME',
            '2-INHERITANCE',
            '3-NOT IN NEED OF MONEY',
            '4-PROPERTY ISSUES',
            '5-UNDISCLOSED BUYERS',
            '5-UNDISCLOSED BUYER',  # legacy singular, kept valid for old data
            '6-OTHERS',
            '7-HARD NO',
            # NU (NOT USABLE)
            '1-GEOGRAPHICAL',
            '2-STRUCTURAL',
            '3-GOVERNMENT',
            # SS/SL (SECURED SALE / SECURED LEASE)
            'WITH RELOC SURVEY',
            # OS/OL (OPEN TO SALE / OPEN TO LEASE)
            'FLIPPED',
        ]}

    # YES/NO status dropdowns
    norm_col = re.sub(r'[\s\-_–—/]+', '', col_upper)
    yes_no_cols = {
        'AVAILABILITYSCHEDULINGWITHTHELO',
        'PENDINGTITLEUNDOCUMENTEDLAND',
        'PENDINGTITLEUNTITLEDLAND',
        'PENDINGTITLEODC',
        'WITHANNOTATIONACTIVEADVERSECLAIM',
        'WITHANNOTATIONTHIRDPARTYENCUMBRANCESMORTGAGE',
        'WITHTENANTISSUE',
        'WITHCLAIMANTISSUE',
        'WITHOVERLAPISSUE',
        'LANDOWNERAIFVERIFICATION',
        'VERIFICATIONOFHEIRSAIFOFTHEDECEASEDROS',
        'HEIRSAIFBASEDABROAD',
        'NONEXTOFKINTOTHERO',
        'AWAITINGDECISIONFROMTHELO',
    }
    if (norm_col in yes_no_cols
            or col_upper in (
                'AVAILABILITY/ SCHEDULING WITH THE LO',
                'PENDING TITLE - UNDOCUMENTED LAND',
                'PENDING TITLE - UNTITLED LAND',
                'PENDING TITLE - ODC',
                'PENDING TITLE – ODC',
                'WITH ANNOTATION - ACTIVE ADVERSE CLAIM',
                'WITH ANNOTATION – ACTIVE ADVERSE CLAIM',
                'WITH ANNOTATION - THIRD-PARTY ENCUMBRANCES/MORTGAGE',
                'WITH ANNOTATION - THIRD-PARTY ENCUMBRANCES/ MORTGAGE',
                'WITH TENANT ISSUE',
                'WITH CLAIMANT ISSUE',
                'WITH OVERLAP ISSUE',
                'LANDOWNER/AIF VERIFICATION',
                'VERIFICATION OF HEIRS/AIF OF THE DECEASED RO/S',
                'HEIRS/AIF BASED ABROAD',
                'NO NEXT OF KIN TO THE RO',
                'AWAITING DECISION FROM THE LO',
            )):
        return {'type': 'choice', 'choices': ['YES', 'NO']}

    if col_upper == 'MUNICODE':
        choices = ['UNASSIGNED']
        for m in municipality_codes:
            code = m.get('muniCode')
            if code and code not in choices:
                choices.append(code)
        return {'type': 'choice', 'choices': choices}

    if col_upper in ('MATCHED NEGOTIATOR NAME', 'MATCHED SOURCER', 'MATCHED SOURCER NAME', 'MATCHED NEGOTIATOR'):
        choices = [t.get('employeeName', '') for t in filter_teams_by_mode(teams, mode) if t.get('employeeName')]
        return {'type': 'choice', 'choices': sorted(list(set(choices)))}


    return {'type': 'text'}



def build_team_lookup(teams: list[dict]) -> dict:
    lookup = {}
    for t in teams:
        name = (t.get('employeeName', '') or '').strip()
        if name:
            lookup[name.lower()] = {'correctTeam': normalize_team(t.get('team', '')),
                                    'correctGroup': normalize_group(t.get('group', ''))}
    return lookup


def negotiator_options(teams: list[dict], mode: str = 'negotiation') -> list[str]:
    scoped = filter_teams_by_mode(teams, mode)
    return sorted({(t.get('employeeName', '') or '').strip() for t in scoped if (t.get('employeeName', '') or '').strip()})


def calculate_review_rows(rows: list[dict], municipality_codes: list[dict], mapping_statuses: list[dict],
                          build_rows: list[dict], existing_records: list[dict],
                          existing_area_set: set[str] | None = None,
                          ignore_build_work_week: bool = False,
                          mode: str = 'negotiation') -> list[dict]:
    """Apply the review-table derivations (runReviewCalculations)."""
    if existing_area_set is None and existing_records:
        existing_area_set = {
            str(r.get('aREAINDEX') or r.get('AREA-INDEX') or r.get('AREA INDEX') or r.get('areaIndex') or '').strip().lower()
            for r in (existing_records or [])
            if (r.get('aREAINDEX') or r.get('AREA-INDEX') or r.get('AREA INDEX') or r.get('areaIndex'))
        }

    area_calculated = []
    for row in rows:
        dn = normalize_csv_date_columns(row)
        derived_ai = get_derived_area_index(dn, municipality_codes)
        dn = {**dn, 'AREA-INDEX': derived_ai, 'AREA INDEX': derived_ai,
              'MUNICODE': get_derived_municode(dn, municipality_codes)}
        area_calculated.append(dn)

    # Pre-index batch by (area_index, row_date) for instant duplicate detection
    batch_by_area_date: dict[tuple[str, str], list[dict]] = collections.defaultdict(list)
    for r in area_calculated:
        ai = str(r.get('AREA INDEX') or '').strip().lower()
        d = get_normalized_record_date(r)
        if ai and d:
            batch_by_area_date[(ai, d)].append(r)

    # Pre-index build rows for instant build lookup
    build_first_by_area: dict[str, str] = {}
    build_by_area_and_ww: dict[tuple[str, int], str] = {}
    for item in (build_rows or []):
        b_ai = str(item.get('areaIndex', '') or '').strip().lower()
        if not b_ai:
            continue
        b_val = str(item.get('build', '') or '').strip()
        if not b_val:
            continue
        if b_ai not in build_first_by_area:
            build_first_by_area[b_ai] = _build_value_or_balance(b_val)
        ww_norm = normalize_work_week(item.get('workWeek'))
        if ww_norm is not None and (b_ai, ww_norm) not in build_by_area_and_ww:
            build_by_area_and_ww[(b_ai, ww_norm)] = _build_value_or_balance(b_val)

    calculated = []
    for row in area_calculated:
        derived_ms = get_derived_mapping_status(row, area_calculated, existing_records, municipality_codes,
                                                existing_area_set=existing_area_set,
                                                batch_by_area_date=batch_by_area_date,
                                                mode=mode)
        row_with_ms = {**row, 'MAPPING STATUS': derived_ms}
        derived_du = get_derived_data_usability(row_with_ms, mapping_statuses, None, existing_records,
                                                municipality_codes, existing_area_set=existing_area_set,
                                                mode=mode)
        calc = {
            **row_with_ms,
            'BUILD': get_derived_build(row, build_rows, ignore_work_week=ignore_build_work_week,
                                       build_first_by_area=build_first_by_area,
                                       build_by_area_and_ww=build_by_area_and_ww),
            'DATA USABILITY': derived_du,
            'YEAR': get_year_from_row_date(row),
            'LOT AREA (HA)': get_lot_area_ha(row.get('LOT AREA (SQM)', '') or ''),
        }
        # Preserve the user's manual selections on derived columns (BUILD, DATA
        # USABILITY, MAPPING STATUS, ...) so re-calculating doesn't revert them.
        overrides = row.get('_overrides')
        if overrides:
            for k, v in overrides.items():
                calc[k] = v
        calculated.append(calc)
    return calculated



def detect_mode(grid: list[list[str]]) -> str:
    headers = {normalize_header(h) for h in (grid[0] if grid else [])}
    canon_headers = {canonical_header(h) for h in (grid[0] if grid else [])}
    if ('LSA NAME' in headers or 'REPORT DATE' in headers or
        'LSANAME' in canon_headers or 'REPORTDATE' in canon_headers):
        if ('NEGOTIATOR NAME' not in headers and 'NEGO DATE' not in headers and
            'NEGOTIATORNAME' not in canon_headers and 'NEGODATE' not in canon_headers):
            return 'land-sourcing'
    return 'negotiation'



def run_pipeline(input_path, output_path, reference_path=None, existing_path=None, mode=None):
    probe_mode = mode or 'negotiation'
    grid = read_grid_from_file(input_path, probe_mode)
    if mode is None:
        mode = detect_mode(grid)
        if mode != probe_mode:
            grid = read_grid_from_file(input_path, mode)

    added_columns = add_missing_upload_columns(grid, mode)
    if added_columns:
        print(f'Note: added {len(added_columns)} missing column(s) as blank: {", ".join(added_columns)}')
    rows = parse_grid(grid, mode)
    if not rows:
        raise ValueError('No data rows found in the input file.')

    # Reference data
    teams, municipality_codes, mapping_statuses, build_rows = [], [], [], []
    if reference_path:
        teams, municipality_codes, mapping_statuses, build_rows = load_reference_workbook(reference_path)

    existing_records = load_existing_records(existing_path) if existing_path else []
    latest_id = max((r['iD1'] for r in existing_records), default=0)
    rows = apply_sequential_ids(rows, latest_id)

    # Calculate derived review values (runReviewCalculations)
    calculated = calculate_review_rows(rows, municipality_codes, mapping_statuses, build_rows, existing_records)

    productivity = build_productivity_rows(calculated, teams, mapping_statuses, existing_records, municipality_codes)
    enriched = enrich_workspace_rows(productivity, teams)
    write_output(enriched, output_path, mode)
    return len(enriched), mode


# ---------------------------------------------------------------------------
# CLI / interactive entry
# ---------------------------------------------------------------------------

def _interactive():
    print('Ground Data Processing Tool (offline)')
    print('-' * 40)
    input_path = input('Input CSV/Excel file: ').strip().strip('"')
    reference_path = input('Reference workbook (optional, Enter to skip): ').strip().strip('"') or None
    existing_path = input('Existing records file (optional, Enter to skip): ').strip().strip('"') or None
    output_path = input('Output file (.csv or .xlsx): ').strip().strip('"')
    mode_in = input('Mode [auto/negotiation/land-sourcing] (Enter=auto): ').strip().lower() or 'auto'
    mode = None if mode_in in ('', 'auto') else mode_in
    try:
        count, used_mode = run_pipeline(input_path, output_path, reference_path, existing_path, mode)
        print(f'\nDone. Wrote {count} productivity rows ({used_mode}) to:\n  {output_path}')
    except Exception as exc:  # noqa: BLE001
        print(f'\nERROR: {exc}')
    input('\nPress Enter to exit...')


def main(argv=None):
    parser = argparse.ArgumentParser(description='Ground Data Processing Tool - offline CSV/Excel productivity generator.')
    parser.add_argument('--input', '-i', help='Input CSV or Excel file (the uploaded rows).')
    parser.add_argument('--reference', '-r', help='Reference workbook: sheets for Team Composition, Municipality Code, Mapping Status, Batangas Build.')
    parser.add_argument('--existing', '-e', help='Existing records CSV/Excel (for ID base and INDEXED PROPERTY UPDATE detection).')
    parser.add_argument('--output', '-o', help='Output file (.csv or .xlsx).')
    parser.add_argument('--mode', '-m', choices=['negotiation', 'land-sourcing'], help='Load type. Auto-detected if omitted.')
    args = parser.parse_args(argv)

    if not args.input or not args.output:
        if len(sys.argv) == 1:
            _interactive()
            return 0
        parser.error('--input and --output are required (or run with no arguments for interactive mode).')

    count, used_mode = run_pipeline(args.input, args.output, args.reference, args.existing, args.mode)
    print(f'Wrote {count} productivity rows ({used_mode}) to {args.output}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
